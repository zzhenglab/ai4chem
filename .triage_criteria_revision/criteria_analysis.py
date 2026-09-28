"""Summarize fixed topic criteria after classification; no ratio-based relabeling."""
import math
from pathlib import Path

import pandas as pd


TOPIC_ORDER = ('Chemical synthesis', 'Theory & modeling', 'Crystal engineering',
               'Functional materials', 'Unclassified')


def yn_ratio(y, n):
    if n:
        return float(y) / n
    return float('inf') if y else float('nan')


def ratio_text(value):
    if pd.isna(value):
        return 'undefined'
    if math.isinf(value):
        return 'infinite (N=0)'
    return '{:.2f}:1'.format(value)


def summarize_criteria(assignments, decisions, criteria_order):
    """One classification per paper per criterion, identical corpus and decisions.

    Thresholds diagnose the resulting classifications. They are not used to
    assign labels, remove papers, select a winner, or update the fixed rules.
    """
    if decisions.paper_id.duplicated().any():
        raise ValueError('Decision input must have one row per paper.')
    if assignments.duplicated(['criterion', 'paper_id']).any():
        raise ValueError('A paper has multiple assignments within a criterion.')
    if set(assignments.criterion) != set(criteria_order):
        raise ValueError('Assignments must contain every declared criterion, and no others.')
    if not set(assignments.category).issubset(set(TOPIC_ORDER)):
        raise ValueError('Unknown topic category.')
    if not set(decisions.triage_decision).issubset({'Y', 'N', 'Unknown', 'Conflict'}):
        raise ValueError('Unsupported decision value; missing decisions must be Unknown.')
    expected_ids = set(decisions.paper_id)
    for criterion in criteria_order:
        if set(assignments.loc[assignments.criterion.eq(criterion), 'paper_id']) != expected_ids:
            raise ValueError('Every criterion must classify the entire same corpus: ' + criterion)
    annotated = assignments.merge(decisions[['paper_id', 'triage_decision']], on='paper_id',
                                  how='left', validate='many_to_one')
    counts, targets = [], []
    for criterion in criteria_order:
        papers = annotated.loc[annotated.criterion.eq(criterion)]
        table = pd.crosstab(papers.category, papers.triage_decision).reindex(
            index=TOPIC_ORDER, columns=['Y', 'N', 'Unknown', 'Conflict'], fill_value=0).astype(int)
        table['Before total'] = table.sum(axis=1)
        table['Y:N ratio'] = [yn_ratio(y, n) for y, n in zip(table.Y, table.N)]
        table['Y:N'] = table['Y:N ratio'].map(ratio_text)
        table['Needs topic review'] = papers.groupby('category').review_needed.sum().reindex(
            TOPIC_ORDER, fill_value=0).astype(int)
        table.index.name = 'Category'
        table.insert(0, 'criterion', criterion)
        counts.append(table.reset_index())
        chem = table.loc['Chemical synthesis']
        crystal = table.loc['Crystal engineering']
        functional = table.loc['Functional materials']
        unclassified = int(table.loc['Unclassified', 'Before total'])
        complete = bool(table[['Unknown', 'Conflict']].to_numpy().sum() == 0)
        chem_pass = bool(chem.Y > 0 and chem.Y >= 4 * chem.N)
        crystal6_pass = bool(crystal.Y > 0 and crystal.Y > 6 * crystal.N)
        crystal7_pass = bool(crystal.Y > 0 and crystal.Y > 7 * crystal.N)
        functional_pass = bool(functional.Y > functional.N)
        total = int(table['Before total'].sum())
        assert total == len(decisions)
        targets.append({
            'criterion': criterion,
            'Chemical Y:N': chem['Y:N'], 'Chemical >=4:1': chem_pass,
            'Crystal Y:N': crystal['Y:N'], 'Crystal >6:1': crystal6_pass,
            'Crystal >7:1': crystal7_pass,
            'Functional Y:N': functional['Y:N'], 'Functional Y>N': functional_pass,
            'All targets (crystal >6)': complete and chem_pass and crystal6_pass and functional_pass,
            'All targets (crystal >7)': complete and chem_pass and crystal7_pass and functional_pass,
            'Total papers': total, 'Unclassified': unclassified,
            'Topic coverage (%)': round(100 * (total - unclassified) / total, 2) if total else float('nan'),
            'Complete decisions': complete,
        })
    return annotated, pd.concat(counts, ignore_index=True), pd.DataFrame(targets)


def plot_criterion_comparison(counts, targets, criteria_order, labels, output_dir):
    """Draw every fixed criterion, including Unclassified, on a common count scale."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import MaxNLocator
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    colors = {'Y': '#8fbd83', 'N': '#efbd6d'}
    max_count = max(1, int(counts[['Y', 'N']].to_numpy().max()))
    panel_titles = [labels[c] for c in criteria_order]

    def draw(ax, criterion, title):
        table = counts.loc[counts.criterion.eq(criterion)].set_index('Category').reindex(TOPIC_ORDER)
        positions = list(range(len(TOPIC_ORDER)))
        for offset, decision in [(-0.17, 'N'), (0.17, 'Y')]:
            bars = ax.barh([p + offset for p in positions], table[decision], height=0.30,
                           color=colors[decision], label=decision)
            ax.bar_label(bars, labels=['{:,}'.format(int(v)) for v in table[decision]],
                         padding=3, fontsize=8)
        ax.set_yticks(positions)
        ax.set_yticklabels(TOPIC_ORDER, fontsize=9)
        ax.invert_yaxis()
        ax.set_xlim(0, max_count * 1.20)
        ax.xaxis.set_major_locator(MaxNLocator(4, integer=True))
        ax.set_xlabel('Unique papers', fontsize=9)
        ax.set_title(title, loc='left', fontsize=11, weight='bold', pad=12)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.grid(axis='x', alpha=0.17)
        ax.set_axisbelow(True)

    fig, axes = plt.subplots(2, 2, figsize=(14, 10.3))
    for ax, criterion, title in zip(axes.flat, criteria_order, panel_titles):
        draw(ax, criterion, title)
    handles, legend_labels = axes.flat[0].get_legend_handles_labels()
    fig.legend(handles, legend_labels, loc='upper right', ncol=2, frameon=False,
               bbox_to_anchor=(0.985, 0.995))
    fig.suptitle('Literature triage: sensitivity to topic definitions', x=0.02, ha='left',
                 fontsize=16, weight='bold')
    fig.text(0.02, 0.027, 'Same papers and recorded decisions in every panel. Rules use paper content, never Y/N.',
             fontsize=9, color='#52606d')
    fig.text(0.02, 0.008, 'Exploratory classifications; Unclassified is retained. Ratio targets do not establish classification accuracy.',
             fontsize=9, color='#52606d')
    fig.tight_layout(rect=[0, 0.055, 1, 0.945], h_pad=3, w_pad=2.4)
    fig.savefig(output_dir / 'all_criteria_Y_N.png', dpi=200, bbox_inches='tight')
    fig.savefig(output_dir / 'all_criteria_Y_N.pdf', bbox_inches='tight')

    for criterion, title in zip(criteria_order, panel_titles):
        single, ax = plt.subplots(figsize=(10.5, 6.5))
        draw(ax, criterion, title)
        ax.legend(loc='upper right', frameon=False, ncol=2)
        target = targets.set_index('criterion').loc[criterion]
        text = 'Y:N: chemical {} | crystal {} | functional {}'.format(
            target['Chemical Y:N'], target['Crystal Y:N'], target['Functional Y:N'])
        single.text(0.02, 0.06, text, fontsize=9, color='#52606d')
        single.text(0.02, 0.025, 'Recorded Y/N decisions; exploratory topic definition; unresolved topics remain visible.',
                    fontsize=9, color='#52606d')
        single.tight_layout(rect=[0, 0.10, 1, 1])
        single.savefig(output_dir / (criterion + '_Y_N.png'), dpi=200, bbox_inches='tight')
        single.savefig(output_dir / (criterion + '_Y_N.pdf'), bbox_inches='tight')
        plt.close(single)
    return fig


def write_comparison_workbook(tables, path):
    """Streaming export keeps large per-paper comparison tables within memory limits."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    for name, table in tables.items():
        sheet = workbook.create_sheet(name[:31])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = 'A1:{}{}'.format(get_column_letter(len(table.columns)), len(table) + 1)
        sheet.append(list(table.columns))
        for values in table.itertuples(index=False, name=None):
            row = []
            for value in values:
                if pd.isna(value):
                    value = None
                elif isinstance(value, float) and math.isinf(value):
                    value = 'infinite' if value > 0 else '-infinite'
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'
                row.append(cell)
            sheet.append(row)
    workbook.save(path)
