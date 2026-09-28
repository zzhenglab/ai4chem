"""Read-only consistency checks for executed comparison artifacts."""
from pathlib import Path
import hashlib
import json
import pandas as pd
import nbformat
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
OUT = HERE / 'outputs' / 'criteria_comparison'
nb = nbformat.read(HERE / 'literature_triage_topics.ipynb', as_version=4)
nbformat.validate(nb)
code_cells = [c for c in nb.cells if c.cell_type == 'code']
assert all(c.execution_count is not None for c in code_cells)
assert not any(o.output_type == 'error' for c in code_cells for o in c.outputs)
assert not any('from topic_criteria import' in c.source or 'from criteria_analysis import' in c.source
               or 'from topic_classifier import' in c.source for c in code_cells)
base_nb = nbformat.read(HERE / 'archive' / 'v3_before_criteria' / 'literature_triage_topics.ipynb', as_version=4)
for current, prior in zip(nb.cells, base_nb.cells):
    if current.source != prior.source:
        assert 'def apply_overrides(' in current.source and 'def apply_overrides(' in prior.source)

assign = pd.read_csv(OUT / 'paper_assignments.csv', keep_default_na=False)
counts = pd.read_csv(OUT / 'counts_by_criterion.csv')
targets = pd.read_csv(OUT / 'target_checks.csv')
baseline = pd.read_csv(HERE / 'outputs' / 'full_decision_audit.csv', keep_default_na=False)
assert len(assign) == 4 * len(baseline)
assert not assign.duplicated(['criterion', 'paper_id']).any()
for criterion, frame in assign.groupby('criterion'):
    assert set(frame.paper_id) == set(baseline.paper_id)
    assert frame.triage_decision.value_counts().to_dict() == baseline.triage_decision.value_counts().to_dict()
    table = counts.loc[counts.criterion.eq(criterion)].set_index('Category')
    assert int(table['Before total'].sum()) == len(baseline)
    assert int(table.Y.sum()) == baseline.triage_decision.eq('Y').sum()
    assert int(table.N.sum()) == baseline.triage_decision.eq('N').sum()
    assert frame.loc[frame.is_review_article, 'category'].eq('Theory & modeling').all()
    assert frame.reason.str.len().gt(0).all()
    for category, row in table.iterrows():
        subset = frame.loc[frame.category.eq(category)]
        assert len(subset) == row['Before total']
        assert subset.triage_decision.eq('Y').sum() == row.Y
        assert subset.triage_decision.eq('N').sum() == row.N
    chem, crystal, functional = [table.loc[c] for c in ['Chemical synthesis', 'Crystal engineering', 'Functional materials']]
    target = targets.loc[targets.criterion.eq(criterion)].iloc[0]
    assert target['Chemical >=4:1'] == bool(chem.Y > 0 and chem.Y >= 4 * chem.N)
    assert target['Crystal >6:1'] == bool(crystal.Y > 0 and crystal.Y > 6 * crystal.N)
    assert target['Crystal >7:1'] == bool(crystal.Y > 0 and crystal.Y > 7 * crystal.N)
    assert target['Functional Y>N'] == bool(functional.Y > functional.N)
    assert (OUT / (criterion + '_Y_N.png')).stat().st_size > 10000
    assert (OUT / (criterion + '_Y_N.pdf')).stat().st_size > 1000
    if criterion == 'baseline_current':
        assert frame.set_index('paper_id').category.to_dict() == baseline.set_index('paper_id').category.to_dict()

manifest = json.loads((OUT / 'comparison_manifest.json').read_text())
for path, expected in manifest['inputs'].items():
    assert hashlib.sha256(Path(path).read_bytes()).hexdigest() == expected
assert manifest['classifier_uses_triage_decisions'] is False
assert manifest['criterion_selected_by_ratio'] is None
criteria_cell = next(c.source for c in code_cells if 'def assign_topic_criteria(' in c.source)
assert hashlib.sha256(criteria_cell.encode('utf-8')).hexdigest() == manifest['executed_criteria_code_sha256']
workbook = load_workbook(OUT / 'topic_criteria_comparison.xlsx', read_only=True, data_only=False)
for name in ['criterion_definitions', 'counts_by_criterion', 'target_checks', 'paper_assignments', 'category_changes', 'transition_counts']:
    csv = pd.read_csv(OUT / (name + '.csv'), keep_default_na=False)
    rows = workbook[name].iter_rows(values_only=True)
    assert list(next(rows)) == list(csv.columns)
    assert sum(1 for _ in rows) == len(csv)
workbook.close()
print('PASS: executed notebook, all four corpora/decisions/reviews, ratio targets, PNG/PDF plots, source hashes, and Excel row counts.')
print(targets.to_string(index=False))
